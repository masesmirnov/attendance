from __future__ import annotations

import asyncio
import io
from dataclasses import dataclass
from typing import Any, Mapping

import openpyxl
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload
from openpyxl.utils import column_index_from_string

EXCEL_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(slots=True, frozen=True)
class RosterEntry:
    row: int
    name: str


class DriveExcelClient:
    def __init__(self, drive: Any):
        self._drive = drive
        self._is_excel: dict[str, bool] = {}

    def owns(self, file_id: str) -> bool:
        if file_id not in self._is_excel:
            meta = self._drive.files().get(fileId=file_id, fields="mimeType", supportsAllDrives=True).execute()
            self._is_excel[file_id] = meta["mimeType"] == EXCEL_MIME_TYPE
        return self._is_excel[file_id]

    def _download(self, file_id: str) -> openpyxl.Workbook:
        data = self._drive.files().get_media(fileId=file_id, supportsAllDrives=True).execute()
        return openpyxl.load_workbook(io.BytesIO(data))

    def get_roster(self, file_id: str, worksheet: str, name_col: str, start_row: int) -> list[RosterEntry]:
        sheet = self._download(file_id)[worksheet]
        col = column_index_from_string(name_col)

        out: list[RosterEntry] = []
        for row in range(start_row, sheet.max_row + 1):
            value = sheet.cell(row, col).value
            name = "" if value is None else str(value).strip()
            if name:
                out.append(RosterEntry(row=row, name=name))
        return out

    def get_column_values(self, file_id: str, worksheet: str, col: str, start_row: int, end_row: int) -> list[Any]:
        sheet = self._download(file_id)[worksheet]
        index = column_index_from_string(col)
        cells = (sheet.cell(row, index).value for row in range(start_row, end_row + 1))
        return ["" if value is None else value for value in cells]

    def set_column_values(
        self, file_id: str, worksheet: str, col: str, start_row: int, end_row: int, values: list[Any]
    ) -> None:
        workbook = self._download(file_id)
        sheet = workbook[worksheet]
        index = column_index_from_string(col)
        for row, value in zip(range(start_row, end_row + 1), values, strict=True):
            sheet.cell(row, index).value = None if value == "" else value

        buffer = io.BytesIO()
        workbook.save(buffer)
        media = MediaIoBaseUpload(buffer, mimetype=EXCEL_MIME_TYPE)
        self._drive.files().update(fileId=file_id, media_body=media, supportsAllDrives=True).execute()


class GoogleSheetsClient:
    def __init__(self, service_account_info: Mapping[str, Any]):
        scopes = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
        creds = Credentials.from_service_account_info(dict(service_account_info), scopes=scopes)
        self._svc = build("sheets", "v4", credentials=creds, cache_discovery=False)
        self._excel = DriveExcelClient(build("drive", "v3", credentials=creds, cache_discovery=False))

    def get_roster(self, sheet_id: str, worksheet: str, name_col: str, start_row: int) -> list[RosterEntry]:
        if self._excel.owns(sheet_id):
            return self._excel.get_roster(sheet_id, worksheet, name_col, start_row)

        rng = f"{worksheet}!{name_col}{start_row}:{name_col}"
        resp = self._svc.spreadsheets().values().get(spreadsheetId=sheet_id, range=rng).execute()
        values = resp.get("values", [])

        out: list[RosterEntry] = []
        for i, row in enumerate(values):
            name = (row[0] if row else "")
            name = str(name).strip()
            if name:
                out.append(RosterEntry(row=start_row + i, name=name))
        return out

    def get_column_values(
        self, sheet_id: str, worksheet: str, col: str, start_row: int, end_row: int
    ) -> list[Any]:
        if self._excel.owns(sheet_id):
            return self._excel.get_column_values(sheet_id, worksheet, col, start_row, end_row)

        rng = f"{worksheet}!{col}{start_row}:{col}{end_row}"
        resp = self._svc.spreadsheets().values().get(spreadsheetId=sheet_id, range=rng).execute()
        values = resp.get("values", [])

        out: list[Any] = []
        for r in range(start_row, end_row + 1):
            idx = r - start_row
            if idx < len(values) and values[idx]:
                out.append(values[idx][0])
            else:
                out.append("")
        return out

    def set_column_values(
        self,
        sheet_id: str,
        worksheet: str,
        col: str,
        start_row: int,
        end_row: int,
        values: list[Any],
    ) -> None:
        expected = end_row - start_row + 1
        if len(values) != expected:
            raise ValueError(f"values length mismatch: expected {expected}, got {len(values)}")

        if self._excel.owns(sheet_id):
            self._excel.set_column_values(sheet_id, worksheet, col, start_row, end_row, values)
            return

        rng = f"{worksheet}!{col}{start_row}:{col}{end_row}"
        body = {"values": [[v] for v in values]}
        self._svc.spreadsheets().values().update(
            spreadsheetId=sheet_id,
            range=rng,
            valueInputOption="RAW",
            body=body,
        ).execute()


class GoogleSheetsService:
    def __init__(self, client: GoogleSheetsClient):
        self._client = client

    async def get_roster(self, sheet_id: str, worksheet: str, name_col: str, start_row: int) -> list[RosterEntry]:
        return await asyncio.to_thread(self._client.get_roster, sheet_id, worksheet, name_col, start_row)

    async def get_column_values(
        self, sheet_id: str, worksheet: str, col: str, start_row: int, end_row: int
    ) -> list[Any]:
        return await asyncio.to_thread(
            self._client.get_column_values, sheet_id, worksheet, col, start_row, end_row
        )

    async def set_column_values(
        self,
        sheet_id: str,
        worksheet: str,
        col: str,
        start_row: int,
        end_row: int,
        values: list[Any],
    ) -> None:
        await asyncio.to_thread(
            self._client.set_column_values, sheet_id, worksheet, col, start_row, end_row, values
        )
